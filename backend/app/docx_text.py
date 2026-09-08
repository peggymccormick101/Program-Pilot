"""Plain-text extraction from an uploaded .docx (the Feature Technical
Specification Document), for feeding into a Claude prompt. Walks the
document body in reading order so headings, paragraphs, and tables (the
FTSD template's "General Information" section is a table) come out in
the same order a person reading the document top-to-bottom would see
them -- iterating doc.paragraphs and doc.tables separately would lose
that order."""

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def extract_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    lines = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, document).text.strip()
            if text:
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
    return "\n".join(lines)
