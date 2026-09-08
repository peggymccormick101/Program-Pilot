"""Renders the structured output of ai.generate_exec_feature_summary() into
the single-slide executive PowerPoint the workflow step promises. Layout,
colors, fonts, and icon set match the reference slide the user approved
(Customer_Profile_Update_Executive_Summary_Updated.pptx) exactly, so every
run looks the same regardless of what Claude returned that time.

Icons are the 5 bundled PNGs extracted from that reference file
(app/assets/icons/) -- reusing the exact assets rather than redrawing them
keeps visual fidelity with what was already approved."""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x1E, 0x27, 0x61)
RUST = RGBColor(0xC9, 0x6A, 0x2E)
CARD_FILL = RGBColor(0xF3, 0xF6, 0xFD)
CARD_BORDER = RGBColor(0xD8, 0xE3, 0xF5)
RISK_FILL = RGBColor(0xFB, 0xEF, 0xE6)
RISK_BORDER = RGBColor(0xEA, 0xD3, 0xBE)
SUBTITLE_BLUE = RGBColor(0xCA, 0xDC, 0xFC)
BODY_TEXT = RGBColor(0x22, 0x26, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

ICONS_DIR = os.path.join(os.path.dirname(__file__), "assets", "icons")

CARD_W, CARD_H = Inches(6.06), Inches(2.28)
CARD_POSITIONS = [
    (Inches(0.5), Inches(1.35)),
    (Inches(6.78), Inches(1.35)),
    (Inches(0.5), Inches(3.79)),
    (Inches(6.78), Inches(3.79)),
]
CARDS = [
    ("business_problem", "Business Problem", "alert-circle.png"),
    ("customer_value", "Customer Value", "heart.png"),
    ("feature_scope", "Feature Scope", "layers.png"),
    ("success_criteria", "Success Criteria", "target.png"),
]


def _no_line(shape):
    shape.line.fill.background()


def _set_bullet(paragraph, *, indent=True):
    """python-pptx has no bullet API -- inject the buChar element the
    reference slide uses directly, matching its exact pPr shape."""
    pPr = paragraph._pPr
    if pPr is None:
        pPr = paragraph._p.get_or_add_pPr()
    if indent:
        pPr.set("marL", "177800")
        pPr.set("indent", "-177800")
    bu_char = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu_char)


def _add_text(slide, left, top, width, height, text, *, size, color, bold=False,
               italic=False, font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color
    return box


def _add_bullets(slide, left, top, width, height, items, *, size, color, font="Calibri", line_spacing=105):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(size * line_spacing / 100)
        run = p.add_run()
        run.text = item
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = color
        _set_bullet(p)
    return box


def _add_rounded_rect(slide, left, top, width, height, fill, *, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border:
        shape.line.color.rgb = border
        shape.line.width = Pt(0.75)
    else:
        _no_line(shape)
    shape.shadow.inherit = False
    return shape


def _add_icon_badge(slide, center_left, center_top, size, icon_file, *, circle_fill):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, center_left, center_top, size, size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = circle_fill
    _no_line(circle)
    circle.shadow.inherit = False
    icon_size = Emu(int(size * 0.58))
    icon_offset = Emu(int((size - icon_size) / 2))
    slide.shapes.add_picture(
        os.path.join(ICONS_DIR, icon_file),
        center_left + icon_offset,
        center_top + icon_offset,
        icon_size,
        icon_size,
    )


def build_exec_summary_pptx(data: dict, *, feature_name: str, priority_label: str | None = None) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    # Header bar
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.15))
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    _no_line(header)
    header.shadow.inherit = False

    _add_text(
        slide, Inches(0.5), Inches(0.14), Inches(9.6), Inches(0.55),
        feature_name, size=26, color=WHITE, bold=True, font="Cambria", anchor=MSO_ANCHOR.BOTTOM,
    )
    _add_text(
        slide, Inches(0.5), Inches(0.68), Inches(9.6), Inches(0.35),
        "Executive Summary  |  Feature Definition", size=13, color=SUBTITLE_BLUE,
        italic=True, anchor=MSO_ANCHOR.MIDDLE,
    )

    if priority_label:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.55), Inches(0.32), Inches(2.28), Inches(0.50)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = RUST
        _no_line(badge)
        badge.shadow.inherit = False
        badge_tf = badge.text_frame
        badge_tf.margin_left = badge_tf.margin_right = badge_tf.margin_top = badge_tf.margin_bottom = 0
        badge_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = badge_tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = priority_label
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.name = "Calibri"
        run.font.color.rgb = WHITE

    # Four quadrant cards
    for (key, title, icon_file), (card_left, card_top) in zip(CARDS, CARD_POSITIONS):
        _add_rounded_rect(slide, card_left, card_top, CARD_W, CARD_H, CARD_FILL, border=CARD_BORDER)
        _add_icon_badge(
            slide, card_left + Inches(0.36), card_top + Inches(0.09), Inches(0.62), icon_file, circle_fill=NAVY
        )
        _add_text(
            slide, card_left + Inches(1.15), card_top + Inches(0.24), Inches(4.56), Inches(0.40),
            title, size=17, color=NAVY, bold=True, font="Cambria", anchor=MSO_ANCHOR.MIDDLE,
        )
        _add_bullets(
            slide, card_left + Inches(0.40), card_top + Inches(0.75), Inches(5.31), Inches(1.36),
            data.get(key) or [], size=11, color=BODY_TEXT,
        )

    # Bottom business-risks bar
    risk_left, risk_top = Inches(0.5), Inches(6.23)
    risk_w, risk_h = Inches(12.33), Inches(0.87)
    _add_rounded_rect(slide, risk_left, risk_top, risk_w, risk_h, RISK_FILL, border=RISK_BORDER)
    _add_icon_badge(
        slide, risk_left + Inches(0.32), risk_top + Inches(0.19), Inches(0.50),
        "warning-triangle.png", circle_fill=RUST,
    )
    _add_text(
        slide, risk_left + Inches(0.95), risk_top + Inches(0.09), Inches(1.85), Inches(0.69),
        "Business Risks", size=14.5, color=NAVY, bold=True, font="Cambria", anchor=MSO_ANCHOR.MIDDLE,
    )

    risks = (data.get("business_risks") or [])[:4]
    risk_slot_positions = [
        (Inches(3.00), Inches(0.08)),
        (Inches(7.52), Inches(0.08)),
        (Inches(3.00), Inches(0.44)),
        (Inches(7.52), Inches(0.44)),
    ]
    for risk_text, (dx, dy) in zip(risks, risk_slot_positions):
        _add_bullets(
            slide, risk_left + dx, risk_top + dy, Inches(4.37), Inches(0.35),
            [risk_text], size=10.5, color=BODY_TEXT, line_spacing=100,
        )

    from io import BytesIO
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()
